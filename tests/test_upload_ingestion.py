import asyncio
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest

from deep_research.contracts.research import UploadSearchOperation
from deep_research.uploads.index import (
    InMemoryUploadIndex,
    OpenSearchUploadIndex,
    UploadIndexSearchAdapter,
)
from deep_research.uploads.ingestion import (
    LocalUploadArtifactStore,
    MalwareDetectedError,
    UploadIngestionService,
    UploadValidationError,
)
from deep_research.uploads.models import IndexedUploadChunk


class CleanScanner:
    async def scan(self, content: bytes) -> None:
        return None


class InfectedScanner:
    async def scan(self, content: bytes) -> None:
        raise MalwareDetectedError("infected")


def make_chunk() -> IndexedUploadChunk:
    return IndexedUploadChunk(
        chunk_id="a" * 64,
        upload_id="upload-1",
        filename="market.md",
        title="Market report",
        content="EV market sales increased substantially in 2025.",
        location="Lines 1-2",
        document_hash="b" * 64,
        access_date="2026-08-27",
        ordinal=0,
    )


@pytest.mark.asyncio
async def test_text_upload_is_quarantined_parsed_stored_and_searchable(tmp_path: Path) -> None:
    index = InMemoryUploadIndex()
    service = UploadIngestionService(
        index=index,
        artifact_store=LocalUploadArtifactStore(
            tmp_path, tenant_id="tenant-a", run_id="run-a"
        ),
        malware_scanner=CleanScanner(),
        chunk_chars=500,
        overlap_chars=50,
    )
    content = (
        b"# Market forecast\n\n"
        b"Internal analysis estimates two million EV sales in 2027.\n"
    )

    result = await service.ingest(
        upload_id="upload-1",
        filename="forecast.md",
        declared_media_type="text/markdown",
        content=content,
        access_date="2026-08-27",
    )
    matches = await UploadIndexSearchAdapter(index).search_uploads(
        UploadSearchOperation(
            query="two million EV sales",
            research_question_ids=["market_size"],
            max_chunks=5,
        )
    )

    assert result.chunk_count == 2
    original = Path(result.original_location)
    extracted = Path(result.extracted_location)
    assert await asyncio.to_thread(original.read_bytes) == content
    assert await asyncio.to_thread(extracted.exists)
    assert matches[0].upload_id == "upload-1"
    assert matches[0].location.startswith("Lines")


@pytest.mark.asyncio
async def test_invalid_upload_is_removed_from_quarantine(tmp_path: Path) -> None:
    store = LocalUploadArtifactStore(tmp_path, tenant_id="tenant-a", run_id="run-a")
    service = UploadIngestionService(
        index=InMemoryUploadIndex(),
        artifact_store=store,
        malware_scanner=CleanScanner(),
    )

    with pytest.raises(UploadValidationError, match="magic bytes"):
        await service.ingest(
            upload_id="upload-2",
            filename="fake.pdf",
            declared_media_type="application/pdf",
            content=b"not a pdf",
            access_date="2026-08-27",
        )

    assert await asyncio.to_thread(lambda: list(tmp_path.rglob("fake.pdf"))) == []


@pytest.mark.asyncio
async def test_malware_rejection_happens_before_parsing_or_indexing(tmp_path: Path) -> None:
    index = InMemoryUploadIndex()
    service = UploadIngestionService(
        index=index,
        artifact_store=LocalUploadArtifactStore(
            tmp_path, tenant_id="tenant-a", run_id="run-a"
        ),
        malware_scanner=InfectedScanner(),
    )

    with pytest.raises(MalwareDetectedError):
        await service.ingest(
            upload_id="upload-3",
            filename="notes.txt",
            declared_media_type="text/plain",
            content=b"malicious content",
            access_date="2026-08-27",
        )

    assert await index.search("malicious", 5) == []
    assert await asyncio.to_thread(lambda: list(tmp_path.rglob("notes.txt"))) == []


@pytest.mark.asyncio
async def test_docx_upload_preserves_paragraph_locations(tmp_path: Path) -> None:
    index = InMemoryUploadIndex()
    service = UploadIngestionService(
        index=index,
        artifact_store=LocalUploadArtifactStore(
            tmp_path, tenant_id="tenant-a", run_id="run-docx"
        ),
        malware_scanner=CleanScanner(),
    )
    content = await asyncio.to_thread(_make_docx)

    result = await service.ingest(
        upload_id="upload-docx",
        filename="brief.docx",
        declared_media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=content,
        access_date="2026-08-27",
    )
    matches = await index.search("charging infrastructure", 5)

    assert result.block_count == 2
    assert matches[0].location.endswith("paragraph 2")


@pytest.mark.asyncio
async def test_pdf_upload_preserves_page_location(tmp_path: Path) -> None:
    index = InMemoryUploadIndex()
    service = UploadIngestionService(
        index=index,
        artifact_store=LocalUploadArtifactStore(
            tmp_path, tenant_id="tenant-a", run_id="run-pdf"
        ),
        malware_scanner=CleanScanner(),
    )

    result = await service.ingest(
        upload_id="upload-pdf",
        filename="market.pdf",
        declared_media_type="application/pdf",
        content=await asyncio.to_thread(_make_pdf),
        access_date="2026-08-27",
    )
    matches = await index.search("two million units", 5)

    assert result.block_count == 1
    assert matches[0].location == "Page 1"


class FakeIndices:
    def __init__(self) -> None:
        self.created: list[dict[str, Any]] = []

    def exists(self, **kwargs: Any) -> bool:
        return False

    def create(self, **kwargs: Any) -> dict[str, bool]:
        self.created.append(kwargs)
        return {"acknowledged": True}


class FakeOpenSearchClient:
    def __init__(self) -> None:
        self.indices = FakeIndices()
        self.bulk_calls: list[dict[str, Any]] = []
        self.search_calls: list[dict[str, Any]] = []
        self.delete_calls: list[dict[str, Any]] = []

    def bulk(self, **kwargs: Any) -> dict[str, bool]:
        self.bulk_calls.append(kwargs)
        return {"errors": False}

    def search(self, **kwargs: Any) -> dict[str, Any]:
        self.search_calls.append(kwargs)
        source = {
            **make_chunk().model_dump(mode="json"),
            "tenant_id": "tenant-a",
            "run_id": "run-a",
        }
        return {"hits": {"hits": [{"_source": source}]}}

    def delete_by_query(self, **kwargs: Any) -> dict[str, int]:
        self.delete_calls.append(kwargs)
        return {"deleted": 1}


@pytest.mark.asyncio
async def test_opensearch_index_injects_and_filters_tenant_and_run_identity() -> None:
    client = FakeOpenSearchClient()
    index = OpenSearchUploadIndex(
        client,
        tenant_id="tenant-a",
        run_id="run-a",
    )

    await index.index_chunks([make_chunk()])
    results = await index.search("EV market", 5)
    await index.delete_upload("upload-1")

    indexed_document = client.bulk_calls[0]["body"][1]
    assert indexed_document["tenant_id"] == "tenant-a"
    assert indexed_document["run_id"] == "run-a"
    filters = client.search_calls[0]["body"]["query"]["bool"]["filter"]
    assert {"term": {"tenant_id": "tenant-a"}} in filters
    assert {"term": {"run_id": "run-a"}} in filters
    assert results == [make_chunk()]


def _make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("EV strategy", level=1)
    document.add_paragraph("Charging infrastructure is the primary constraint.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _make_pdf() -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (EV sales reached two million units.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()
